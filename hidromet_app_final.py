
import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from io import BytesIO
from docx import Document
from fpdf import FPDF
from datetime import datetime
import base64

st.set_page_config(layout="wide")

# Estado de sesión
if 'authenticated' not in st.session_state:
    st.session_state.authenticated = False

# Función para exportar a PDF
def generar_pdf(df):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for i, row in df.iterrows():
        line = ', '.join(str(x) for x in row)
        pdf.cell(200, 10, txt=line, ln=True)
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

# Función para exportar a Word
def generar_word(df):
    doc = Document()
    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def login():
    st.title("Panel de Control Hidrometeorológico")
    username = st.text_input("Usuario")
    password = st.text_input("Contraseña", type="password")
    if st.button("Iniciar sesión"):
        if username == "admin" and password == "admin123":
            st.session_state.authenticated = True
            st.rerun()
        else:
            st.error("Credenciales inválidas")

def logout():
    st.session_state.authenticated = False
    st.rerun()

def admin_panel():
    st.sidebar.title("Panel de Administración")
    if st.sidebar.button("Cerrar sesión"):
        logout()

    uploaded_file = st.sidebar.file_uploader("Subir archivo CSV", type=["csv"])
    if uploaded_file:
        df = pd.read_csv(uploaded_file)
        st.write("Vista previa de los datos:")
        st.dataframe(df)

        with st.expander("Análisis Estadístico"):
            st.write(df.describe())

        with st.expander("Gráficos Interactivos"):
            for col in df.select_dtypes(include=[np.number]).columns:
                st.subheader(f"Histograma de {col}")
                fig = px.histogram(df, x=col)
                st.plotly_chart(fig, use_container_width=True)

        with st.expander("Mapa de Calor de Correlación"):
            st.write(px.imshow(df.corr(), text_auto=True))

        with st.expander("Exportar Informe"):
            col1, col2 = st.columns(2)
            with col1:
                pdf_buffer = generar_pdf(df)
                b64_pdf = base64.b64encode(pdf_buffer.read()).decode()
                href = f'<a href="data:application/pdf;base64,{b64_pdf}" download="informe.pdf">Descargar PDF</a>'
                st.markdown(href, unsafe_allow_html=True)

            with col2:
                word_buffer = generar_word(df)
                b64_word = base64.b64encode(word_buffer.read()).decode()
                href = f'<a href="data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64,{b64_word}" download="informe.docx">Descargar Word</a>'
                st.markdown(href, unsafe_allow_html=True)

    else:
        st.warning("Por favor suba un archivo para acceder a las herramientas.")

def main():
    if st.session_state.authenticated:
        admin_panel()
    else:
        login()

if __name__ == "__main__":
    main()
