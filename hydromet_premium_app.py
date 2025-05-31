
import streamlit as st
import pandas as pd
import plotly.express as px
from fpdf import FPDF
from io import BytesIO
from docx import Document

# ----------------- Estilos Premium Modernos -----------------
st.markdown("""
    <style>
    html, body, [class*="css"] {
        font-family: 'Segoe UI', sans-serif;
        background: linear-gradient(to bottom, #1f1c2c, #928dab);
        color: #ffffff;
    }
    .stApp {
        background-color: #121212;
        border-radius: 10px;
        padding: 20px;
    }
    .stButton > button {
        background-color: #6a11cb;
        background-image: linear-gradient(315deg, #6a11cb 0%, #2575fc 74%);
        color: white;
        border-radius: 8px;
        padding: 0.6em 1.2em;
        border: none;
    }
    .stTextInput > div > div > input {
        background-color: #2c2f33;
        color: #ffffff;
        border-radius: 5px;
        padding: 0.5em;
    }
    .block-container {
        padding-top: 2rem;
    }
    </style>
""", unsafe_allow_html=True)

# ----------------- Usuarios -----------------
USUARIOS = {
    "admin": "YZ1BKzgHIK7P7ZrB",
    "tecnico": "tecnico123"
}

# ----------------- PDF -----------------
def generar_pdf(df):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Reporte de Datos", ln=True, align="C")
    pdf.ln()
    col_width = pdf.w / (len(df.columns) + 1)
    for col in df.columns:
        pdf.cell(col_width, 10, str(col), border=1)
    pdf.ln()
    for row in df.itertuples(index=False):
        for item in row:
            pdf.cell(col_width, 10, str(item), border=1)
        pdf.ln()
    return pdf.output(dest='S').encode('latin-1')

# ----------------- Word -----------------
def generar_word(df):
    doc = Document()
    doc.add_heading("Reporte de Datos", 0)
    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = col
    for index, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, col_name in enumerate(df.columns):
            row_cells[i].text = str(row[col_name])
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# ----------------- Cargar CSV -----------------
def cargar_datos():
    uploaded_file = st.file_uploader("📁 Cargar archivo CSV", type=["csv"])
    if uploaded_file:
        try:
            df = pd.read_csv(uploaded_file)
            if 'fecha' in df.columns:
                df['fecha'] = pd.to_datetime(df['fecha'])
                df.set_index('fecha', inplace=True)
                st.success("✅ CSV cargado con 'fecha' como índice.")
            st.session_state.df_cargado = df
        except Exception as e:
            st.error(f"Error al cargar el archivo: {e}")
            st.session_state.df_cargado = None

# ----------------- Herramientas Básicas -----------------
def herramientas_basicas(df):
    st.subheader("📄 Vista Previa")
    st.dataframe(df)
    st.metric("Total de registros", len(df))
    st.subheader("📈 Serie de Tiempo")
    st.line_chart(df)
    numeric_df = df.select_dtypes(include='number')
    if not numeric_df.empty:
        st.subheader("📊 Matriz de Correlación")
        fig = px.imshow(numeric_df.corr(), text_auto=True)
        st.plotly_chart(fig)
    st.subheader("📤 Exportar")
    st.download_button("📄 PDF", data=generar_pdf(df), file_name="reporte.pdf", mime="application/pdf")
    st.download_button("📝 Word", data=generar_word(df), file_name="reporte.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

# ----------------- Herramientas Avanzadas -----------------
def herramientas_admin_extras(df):
    numeric_df = df.select_dtypes(include='number')
    if not numeric_df.empty:
        st.subheader("📌 Histograma Avanzado")
        col_hist = st.selectbox("Columna", numeric_df.columns)
        fig_hist = px.histogram(df, x=col_hist, color_discrete_sequence=["#ff6a00"])
        st.plotly_chart(fig_hist)
        st.subheader("📌 Boxplot")
        fig_box = px.box(df, y=col_hist)
        st.plotly_chart(fig_box)
        st.subheader("📉 Análisis de Tendencia")
        df_rolling = df[numeric_df.columns].rolling(window=7).mean()
        st.line_chart(df_rolling)
        st.subheader("🧮 Nulos por Columna")
        st.write(df.isnull().sum())
        st.subheader("📅 Distribución Mensual")
        st.line_chart(df.resample('M').mean(numeric_only=True))
        st.subheader("📊 Estadísticas Descriptivas")
        st.dataframe(numeric_df.describe())

# ----------------- Panel Técnico -----------------
def tecnico_panel():
    st.title("🔧 Panel Técnico")
    cargar_datos()
    df = st.session_state.get('df_cargado')
    if df is not None:
        herramientas_basicas(df)
    if st.button("Cerrar sesión"):
        cerrar_sesion()

# ----------------- Panel Admin -----------------
def admin_panel():
    st.title("🛠️ Panel de Administración")
    cargar_datos()
    df = st.session_state.get('df_cargado')
    if df is not None:
        herramientas_basicas(df)
        herramientas_admin_extras(df)
    if st.button("Cerrar sesión"):
        cerrar_sesion()

# ----------------- Login -----------------
def login():
    st.title("💧 Hydromet - Centro de Control")
    st.markdown("Inicia sesión para acceder al sistema:")
    usuario = st.text_input("Usuario")
    contraseña = st.text_input("Contraseña", type="password")
    if st.button("Iniciar sesión"):
        if usuario in USUARIOS and USUARIOS[usuario] == contraseña:
            st.session_state.autenticado = True
            st.session_state.usuario = usuario
        else:
            st.error("❌ Usuario o contraseña incorrectos")

# ----------------- Cerrar Sesión -----------------
def cerrar_sesion():
    for key in list(st.session_state.keys()):
        del st.session_state[key]
    st.experimental_rerun()

# ----------------- Inicialización -----------------
if 'autenticado' not in st.session_state:
    st.session_state.autenticado = False

# ----------------- Main -----------------
def main():
    if st.session_state.autenticado:
        if st.session_state.usuario == "admin":
            admin_panel()
        elif st.session_state.usuario == "tecnico":
            tecnico_panel()
    else:
        login()

main()
